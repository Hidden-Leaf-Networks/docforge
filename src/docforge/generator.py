"""High-level document generation API.

Wraps PDF and Word generation with a single, convenient interface.
"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any

from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, Spacer

from docforge.pdf.components import (
    add_metadata_table,
    add_paragraph,
    accent_divider,
    cover_page,
    parse_markdown_content,
    section_header,
    subsection_header,
)
from docforge.pdf.document import ForgeDocument, SimpleForgeDocument
from docforge.theme import DEFAULT_THEME, Theme, build_styles


# Document type labels
DOC_TYPE_LABELS = {
    "business_plan": "BUSINESS PLAN",
    "grant_application": "GRANT APPLICATION",
    "marketing_plan": "MARKETING PLAN",
    "financial_projection": "FINANCIAL PROJECTION",
    "executive_summary": "EXECUTIVE SUMMARY",
    "pitch_deck": "PITCH DECK",
    "report": "REPORT",
    "proposal": "PROPOSAL",
    "whitepaper": "WHITEPAPER",
}


class DocumentGenerator:
    """High-level document generation with configurable theming.

    Args:
        theme: Theme for styling. Uses DEFAULT_THEME if not provided.
    """

    def __init__(self, theme: Theme | None = None):
        self.theme = theme or DEFAULT_THEME
        self.styles = build_styles(self.theme)

    def create_pdf(
        self,
        filepath: str,
        title: str,
        content: dict[str, Any] | str,
        metadata: dict[str, str] | None = None,
    ) -> str:
        """Create a professionally formatted PDF document.

        Args:
            filepath: Output file path.
            title: Document title.
            content: Either a dict with "sections" or "generated_content" keys,
                     or a plain markdown string.
            metadata: Optional dict with author, organization, document_type, location.

        Returns:
            The filepath written to.
        """
        meta = metadata or {}
        doc = ForgeDocument(filepath, theme=self.theme)
        story: list = []

        # Resolve document type label
        doc_type_key = meta.get("document_type", "").lower().replace(" ", "_")
        doc_type = DOC_TYPE_LABELS.get(doc_type_key, meta.get("document_type", "Document").upper())

        cover_page(
            story=story,
            styles=self.styles,
            title=title,
            doc_type=doc_type,
            date_str=datetime.now().strftime("%B %d, %Y"),
            author=meta.get("author", self.theme.brand_name),
            organization=meta.get("organization", ""),
            location=meta.get("location", ""),
            tagline=self.theme.footer_text,
            theme=self.theme,
            client=meta.get("client", ""),
        )

        self._add_content(story, content)
        doc.build(story)
        return filepath

    def create_simple_pdf(
        self,
        filepath: str,
        title: str,
        content: str,
        metadata: dict[str, str] | None = None,
    ) -> str:
        """Create a simpler PDF without cover page distinction.

        Args:
            filepath: Output file path.
            title: Document title.
            content: Markdown or plain text content.
            metadata: Optional dict with author, organization.

        Returns:
            The filepath written to.
        """
        meta = metadata or {}
        doc = SimpleForgeDocument(filepath, theme=self.theme)
        story: list = []

        story.append(Paragraph(title, self.styles["H1"]))
        story.append(accent_divider(self.theme, thickness=2, space_before=4, space_after=16))

        meta_pairs = [
            ["Prepared by:", meta.get("author", self.theme.brand_name)],
            ["Organization:", meta.get("organization", "")],
            ["Date:", datetime.now().strftime("%B %d, %Y")],
        ]
        if meta.get("location"):
            meta_pairs.append(["Location:", meta["location"]])

        add_metadata_table(story, meta_pairs, theme=self.theme)
        story.append(Spacer(1, 0.3 * inch))

        if content:
            parse_markdown_content(content, story, self.styles, theme=self.theme)

        doc.build(story)
        return filepath

    def create_word(
        self,
        filepath: str,
        title: str,
        content: dict[str, Any] | str,
        metadata: dict[str, str] | None = None,
    ) -> str:
        """Create a formatted Word document.

        Requires the `word` extra: `pip install docforge[word]`

        Args:
            filepath: Output file path.
            title: Document title.
            content: Either a dict with "sections" key, or a plain string.
            metadata: Optional dict with author, organization, location.

        Returns:
            The filepath written to.
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.style import WD_STYLE_TYPE
        except ImportError:
            raise ImportError(
                "python-docx is required for Word generation. "
                "Install with: pip install docforge[word]"
            )

        meta = metadata or {}
        doc = Document()

        # Header
        header = doc.sections[0].header
        header.paragraphs[0].text = self.theme.brand_name

        # Title
        title_para = doc.add_heading(title, 0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata table
        meta_rows = [
            ["Created by:", meta.get("author", self.theme.brand_name)],
            ["Organization:", meta.get("organization", "")],
            ["Date:", datetime.now().strftime("%B %d, %Y")],
        ]
        if meta.get("location"):
            meta_rows.append(["Location:", meta["location"]])

        doc.add_paragraph()
        table = doc.add_table(rows=len(meta_rows), cols=2)
        table.style = "Light Grid Accent 1"
        for i, (label, value) in enumerate(meta_rows):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = value

        doc.add_paragraph()
        doc.add_paragraph("_" * 80)
        doc.add_paragraph()

        # Content
        if isinstance(content, dict):
            if "sections" in content:
                for section_title, section_content in content["sections"].items():
                    doc.add_heading(section_title.replace("_", " ").title(), level=1)
                    for paragraph in section_content.split("\n\n"):
                        if paragraph.strip():
                            doc.add_paragraph(paragraph.strip())
                    doc.add_paragraph()
            elif "generated_content" in content:
                self._add_markdown_to_word(doc, content["generated_content"])
        elif isinstance(content, str):
            self._add_markdown_to_word(doc, content)

        # Footer
        footer = doc.sections[0].footer
        footer.paragraphs[0].text = f"{self.theme.footer_text} | {datetime.now().strftime('%B %d, %Y')}"
        footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.save(filepath)
        return filepath

    def _add_content(self, story: list, content: dict[str, Any] | str) -> None:
        """Add content to PDF story from various input formats."""
        if isinstance(content, dict):
            if "sections" in content:
                first = True
                for section_title, section_content in content["sections"].items():
                    formatted = section_title.replace("_", " ").title()
                    section_header(story, self.styles, formatted, new_page=not first, theme=self.theme)
                    first = False
                    if section_content:
                        parse_markdown_content(section_content, story, self.styles, theme=self.theme)
            elif "generated_content" in content:
                generated = content["generated_content"]
                if generated:
                    parse_markdown_content(generated, story, self.styles, theme=self.theme)
            else:
                for key, value in content.items():
                    if key not in ("form_data", "metadata") and value:
                        section_header(story, self.styles, key.replace("_", " ").title(), new_page=False, theme=self.theme)
                        if isinstance(value, str):
                            parse_markdown_content(value, story, self.styles, theme=self.theme)
                        elif isinstance(value, dict):
                            for sub_key, sub_value in value.items():
                                if sub_value:
                                    subsection_header(story, self.styles, sub_key.replace("_", " ").title())
                                    add_paragraph(story, self.styles, str(sub_value))
        elif isinstance(content, str):
            parse_markdown_content(content, story, self.styles, theme=self.theme)

    @staticmethod
    def _add_markdown_to_word(doc, content: str) -> None:
        """Add markdown content to a Word document."""
        import re

        for line in content.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("#"):
                level = len(line) - len(line.lstrip("#"))
                doc.add_heading(line.lstrip("#").strip(), level=min(level, 3))
            elif line.isupper() and len(line) < 60 and line.endswith(":"):
                doc.add_heading(line.rstrip(":"), level=2)
            else:
                doc.add_paragraph(line)
